import unittest
from datetime import date
from io import BytesIO

from openpyxl import Workbook

from medaudit.documents import Document, ElementKind
from medaudit.parsing import XLSXParser


class XLSXParserTest(unittest.TestCase):
    def test_extracts_non_empty_rows_and_sheet_provenance(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        assert worksheet is not None
        worksheet.title = "Synthetic Rules"
        worksheet.append(["Code", "Limit"])
        worksheet.append(["PX-101", 2])
        worksheet.append(["   ", "\n"])
        worksheet.append([])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        document = Document("xlsx-v1", "Synthetic XLSX", "1", date(2026, 1, 1))

        parsed = XLSXParser().parse(document, buffer.getvalue())

        self.assertEqual(len(parsed.elements), 2)
        self.assertEqual(parsed.elements[0].kind, ElementKind.TABLE)
        self.assertEqual(parsed.elements[1].text, "PX-101\t2")
        self.assertEqual(parsed.elements[1].section, "Synthetic Rules")
        self.assertEqual(parsed.elements[1].metadata["row"], "2")

    def test_rejects_invalid_xlsx(self) -> None:
        document = Document("xlsx-v1", "Synthetic", "1", date(2026, 1, 1))

        with self.assertRaisesRegex(ValueError, "invalid or unsupported XLSX"):
            XLSXParser().parse(document, b"not a workbook")
