"""Local XLSX row parser backed by openpyxl."""

import hashlib
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook

from medaudit.documents import Document, DocumentElement, ElementKind
from medaudit.parsing.models import ParsedDocument


class XLSXParser:
    """Represent each non-empty worksheet row as an atomic table element."""

    supported_media_types = frozenset(
        {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    )

    def parse(self, document: Document, content: bytes) -> ParsedDocument:
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
        except Exception as error:
            raise ValueError("invalid or unsupported XLSX content") from error

        elements: list[DocumentElement] = []
        try:
            ordinal = 0
            for worksheet in workbook.worksheets:
                for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                    values = [_render_cell(cell.value) for cell in row]
                    while values and not values[-1]:
                        values.pop()
                    if not any(values):
                        continue
                    text = "\t".join(values)
                    elements.append(
                        DocumentElement(
                            element_id=_element_id(
                                document.document_id,
                                worksheet.title,
                                row_number,
                                text,
                            ),
                            document_id=document.document_id,
                            ordinal=ordinal,
                            kind=ElementKind.TABLE,
                            text=text,
                            section=worksheet.title,
                            metadata={
                                "sheet": worksheet.title,
                                "row": str(row_number),
                                "parser": "openpyxl-rows-v1",
                            },
                        )
                    )
                    ordinal += 1
        finally:
            workbook.close()
        return ParsedDocument(document=document, elements=tuple(elements))


def _render_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return (
        str(value)
        .replace("\t", " ")
        .replace("\r", " ")
        .replace("\n", " ")
        .strip()
    )


def _element_id(document_id: str, sheet: str, row: int, text: str) -> str:
    value = f"{document_id}\0{sheet}\0{row}\0{text}".encode()
    return f"el-{hashlib.sha256(value).hexdigest()[:16]}"
